import base64
import json
from io import BytesIO, StringIO
from typing import Dict, List, Optional, Tuple

import httpx

from src.config import Settings
from src.utils.logging_setup import get_logger

logger = get_logger(__name__)

# 可选依赖
try:
    import pypdf as PyPDF2  # PyPDF2 已弃用，pypdf 为其官方继任者（API 兼容）
except ImportError:
    PyPDF2 = None
try:
    import docx
except ImportError:
    docx = None
try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import csv
except ImportError:
    csv = None


class FileParser:
    def __init__(self, config: Settings):
        self.config = config
        self._client: Optional[httpx.AsyncClient] = None  # 复用的 HTTP 客户端（懒创建）

    def _get_client(self, timeout: float = 30) -> httpx.AsyncClient:
        """复用一个 AsyncClient，避免每次请求都新建连接（省连接开销）"""
        if self._client is None or self._client.is_closed:
            self._client = httpx.AsyncClient(timeout=timeout)
        return self._client

    async def close(self) -> None:
        """优雅关闭：释放复用 HTTP 客户端（修复连接池泄漏）。"""
        if self._client is not None and not self._client.is_closed:
            await self._client.aclose()
            self._client = None

    def _cap(self, text: str) -> str:
        """P1-4 资源上限：提取文本统一截断到 MAX_FILE_TEXT_CHARS，防止超长内容烧 token。"""
        if not text:
            return ""
        limit = getattr(self.config, "MAX_FILE_TEXT_CHARS", 8000)
        if len(text) > limit:
            return text[:limit] + "\n...(内容过长已截断)"
        return text

    # ========== 新增：通过 NapCat HTTP API 获取并解析文件 ==========
    async def fetch_and_parse_file(self, file_id: str, file_name: str) -> Tuple[str, bool]:
        """
        调用 NapCat /get_file 接口获取文件内容，并调用 decode_napcat_file_response 解析
        返回: (提取的文本内容, 是否成功)
        """
        if not file_id:
            return "", False

        try:
            client = self._get_client(timeout=30)
            # 流式读取 + 字节兜底上限：/get_file 返回的是 base64 文本（≈原始字节 × 4/3），
            # 超上限立刻中止，防止 NapCat 返回超大内容时整包载入内存
            max_bytes = int(getattr(self.config, "MAX_FILE_DOWNLOAD_BYTES", 2 * 1024 * 1024))
            cap = max_bytes * 2 + 4096  # base64 上限 + 余量
            body = b""
            rejected = False
            async with client.stream(
                "GET",
                f"{self.config.HTTP_API_BASE}/get_file",
                params={"file_id": file_id},
            ) as resp:
                if resp.status_code != 200:
                    logger.error(f"Fetch file {file_id} failed: HTTP {resp.status_code}")
                    return "", False
                cl = resp.headers.get("content-length")
                if cl and cl.isdigit() and int(cl) > cap:
                    rejected = True
                else:
                    async for chunk in resp.aiter_bytes():
                        body += chunk
                        if len(body) > cap:
                            rejected = True
                            body = b""
                            break
            if rejected:
                logger.error(f"Fetch file {file_id} response exceeds limit ({cap} bytes), aborted")
                return "", False

            # 调用已有的解码方法
            return self.decode_napcat_file_response(body.decode("utf-8", errors="ignore"), file_name)

        except httpx.TimeoutException:
            logger.error(f"Fetch file {file_id} timeout")
            return "", False
        except httpx.HTTPError as e:
            logger.error(f"Fetch file {file_id} HTTP error: {e}")
            return "", False
        except Exception as e:
            logger.exception(f"Fetch file {file_id} unexpected error: {e}")
            return "", False

    # ========== 解码 NapCat 文件响应 ==========
    def decode_napcat_file_response(self, response_text: str, file_name: str) -> Tuple[str, bool]:
        """从 /get_file 返回的 JSON 中提取 base64 并解码"""
        try:
            data = json.loads(response_text)
            if data.get("retcode") != 0:
                return "", False
            file_data = data.get("data", {})
            b64 = file_data.get("base64", "")
            if not b64:
                return "", False
            content_bytes = base64.b64decode(b64)
            # 兜底字节上限（防 NapCat 返回超预期内容，不信任上传通知里的 file_size）
            max_bytes = getattr(self.config, "MAX_FILE_DOWNLOAD_BYTES", 2 * 1024 * 1024)
            if len(content_bytes) > max_bytes:
                logger.error(f"File decode bytes exceed limit: {len(content_bytes)} > {max_bytes}")
                return "", False
            ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
            extracted_text = ""

            if ext == 'txt' or ext == '':
                for enc in ['utf-8', 'gbk']:
                    try:
                        extracted_text = content_bytes.decode(enc)
                        return self._cap(extracted_text), True
                    except UnicodeDecodeError:
                        continue
                return self._cap(content_bytes.decode('utf-8', errors='ignore')), True

            elif ext == 'pdf':
                if PyPDF2 is None:
                    return "", False
                try:
                    pdf_reader = PyPDF2.PdfReader(BytesIO(content_bytes))
                    pages = [page.extract_text() or "" for page in pdf_reader.pages[:self.config.MAX_PDF_PAGES]]
                    extracted_text = self._cap("\n".join(pages))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"PDF parse error: {e}")
                    return "", False

            elif ext == 'docx':
                if docx is None:
                    return "", False
                try:
                    doc = docx.Document(BytesIO(content_bytes))
                    paragraphs = [p.text for p in doc.paragraphs]
                    extracted_text = self._cap("\n".join(paragraphs))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"DOCX parse error: {e}")
                    return "", False

            elif ext == 'xlsx':
                if openpyxl is None:
                    return "", False
                try:
                    # zip 炸弹预检：综合解压尺寸上限（xlsx 是 zip，解压可膨胀数百倍）
                    import zipfile as _zip
                    try:
                        with _zip.ZipFile(BytesIO(content_bytes)) as zf:
                            total_uncompressed = sum(i.file_size for i in zf.infolist())
                            if total_uncompressed > 512 * 1024 * 1024:
                                logger.warning("xlsx decompression bomb: %s", total_uncompressed)
                                return "", False
                    except _zip.BadZipFile:
                        return "", False
                    wb = openpyxl.load_workbook(BytesIO(content_bytes), data_only=True)
                    rows = []
                    cell_count = 0
                    for sheet in wb.worksheets:
                        rows.append(f"--- 工作表: {sheet.title} ---")
                        for row in sheet.iter_rows(values=True):
                            rows.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
                            cell_count += len(row)
                            if cell_count >= self.config.MAX_EXCEL_CELLS:
                                rows.append("...(已达到解析上限)")
                                break
                        if cell_count >= self.config.MAX_EXCEL_CELLS:
                            break
                    extracted_text = self._cap("\n".join(rows))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"XLSX parse error: {e}")
                    return "", False

            elif ext == 'csv':
                if csv is None or StringIO is None:
                    return "", False
                try:
                    for enc in ['utf-8', 'gbk', 'gb2312']:
                        try:
                            text_content = content_bytes.decode(enc)
                            break
                        except UnicodeDecodeError:
                            continue
                    else:
                        text_content = content_bytes.decode('utf-8', errors='ignore')
                    reader = csv.reader(StringIO(text_content))
                    rows = [",".join(row) for row in reader][:self.config.MAX_CSV_ROWS]
                    extracted_text = self._cap("\n".join(rows))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"CSV parse error: {e}")
                    return "", False
            else:
                # 尝试作为文本
                try:
                    extracted_text = self._cap(content_bytes.decode('utf-8'))
                    return extracted_text, True
                except Exception:
                    return "", False
        except json.JSONDecodeError:
            # 非 JSON 响应（如 NapCat 返回错误页）一律视为失败，绝不把原始响应当文件内容
            logger.error("NapCat get_file 响应不是合法 JSON，已拒绝")
            return "", False
        except Exception as e:
            logger.error(f"File decode error: {e}")
            return "", False

    # ========== 提取合并转发消息 ==========
    async def extract_forward_messages(self, message_array: List[Dict]) -> Tuple[str, List[str], bool]:
        """提取合并转发消息中的文本与图片 URL。

        返回: (转发文本, 转发内的图片url列表, 是否有转发)
        图片 url 由调用方逐个交给视觉模型识图（让花璃看到转发里的每一张图）。
        """

        def extract_all_text(obj, sender="未知", prefix="", urls=None):
            results = []
            if isinstance(obj, dict):
                sender_val = obj.get("sender", {})
                if isinstance(sender_val, dict):
                    sender = sender_val.get("user_id", sender)
                elif isinstance(sender_val, (int, str)):
                    sender = str(sender_val)
                # 收集 image 段的 url（NapCat 合并转发里的图片同样带 url）
                if obj.get("type") == "image":
                    img_url = (obj.get("data") or {}).get("url", "")
                    if img_url and urls is not None and img_url not in urls:
                        urls.append(img_url)
                if "text" in obj and isinstance(obj["text"], str):
                    if obj["text"].strip():
                        results.append(f"[用户{sender}]：{obj['text']}")
                for key, value in obj.items():
                    if key == "text":
                        continue
                    results.extend(extract_all_text(value, sender, prefix + key + ".", urls))
            elif isinstance(obj, list):
                for idx, item in enumerate(obj):
                    results.extend(extract_all_text(item, sender, prefix + f"[{idx}]", urls))
            return results

        async def fetch_forward_messages(forward_id) -> Optional[List[Dict]]:
            """通过 NapCat HTTP API 拉取一次转发内容（带缓存去重，同一 id 只拉一次）。"""
            if forward_id in seen_forwards:
                logger.debug(f"Forward {forward_id} served from cache")
                return seen_forwards[forward_id]
            if state["fetches"] >= max_fetches:
                logger.warning(f"Forward fetch budget exceeded ({max_fetches}), stop fetching")
                return None
            state["fetches"] += 1
            try:
                client = self._get_client(timeout=10)
                resp = await client.get(
                    f"{self.config.HTTP_API_BASE}/get_forward_msg",
                    params={"message_id": forward_id}
                )
                if resp.status_code == 200:
                    result = resp.json()
                    if result.get("retcode") == 0:
                        messages = result.get("data", {}).get("messages")
                        seen_forwards[forward_id] = messages
                        return messages
            except Exception as e:
                logger.error(f"Get forward msg error: {e}")
            return None

        async def resolve_nested_forwards(node, depth: int = 0):
            """递归展开嵌套转发（带缓存 + 节点/消息/拉取预算，防套娃转发 DoS）。

            - MAX_FORWARD_DEPTH：展开深度上限
            - MAX_FORWARD_NODES：递归遍历节点总数上限
            - MAX_FORWARD_MESSAGES：展开后消息总数上限（超出截断）
            - MAX_FORWARD_FETCHES：/get_forward_msg 拉取次数上限
            - seen_forwards：同一 forward id 只拉一次（防重复请求）
            """
            state["nodes"] += 1
            if depth > max_depth or state["nodes"] > max_nodes:
                return node
            if isinstance(node, list):
                # 消息总数预算：超出直接截断
                if state["messages"] + len(node) > max_messages:
                    node = node[: max(0, max_messages - state["messages"])]
                    logger.warning(f"Forward message budget exceeded ({max_messages}), truncated")
                state["messages"] += len(node)
                return [await resolve_nested_forwards(item, depth) for item in node]
            if not isinstance(node, dict):
                return node
            if node.get("type") == "forward":
                data = node.get("data") or {}
                messages = data.get("messages")
                if not messages:
                    fid = data.get("id")
                    if fid:
                        fetched = await fetch_forward_messages(fid)
                        if fetched:
                            data["messages"] = fetched
                            messages = fetched
                if messages:
                    return await resolve_nested_forwards(messages, depth + 1)
                return node
            # 普通节点：递归展开所有字段（含 message 数组里的嵌套 forward）
            return {key: await resolve_nested_forwards(value, depth) for key, value in node.items()}

        # 转发预算状态（一次 extract_forward_messages 调用共享）
        max_depth = getattr(self.config, "MAX_FORWARD_DEPTH", 5)
        max_messages = getattr(self.config, "MAX_FORWARD_MESSAGES", 100)
        max_nodes = getattr(self.config, "MAX_FORWARD_NODES", 500)
        max_fetches = getattr(self.config, "MAX_FORWARD_FETCHES", 20)
        state = {"nodes": 0, "messages": 0, "fetches": 0}
        seen_forwards: Dict[str, List[Dict]] = {}

        if not isinstance(message_array, list):
            logger.warning("message_array 不是列表，跳过转发解析")
            return "", [], False

        for msg in message_array:
            if msg.get("type") == "forward":
                forward_data = msg.get("data", {})
                messages = forward_data.get("messages")
                if not messages:
                    messages = await fetch_forward_messages(forward_data.get("id"))
                    if not messages:
                        continue
                # 先展开嵌套转发，再一次性提取文本与图片 url
                resolved = await resolve_nested_forwards(messages)
                image_urls: List[str] = []
                text_lines = extract_all_text(resolved, urls=image_urls)
                if text_lines or image_urls:
                    return "\n".join(text_lines), image_urls, True
                return "", image_urls, False
        return "", [], False

    # ========== 提取 JSON 卡片内容 ==========
    def extract_json_card_content(self, message_array: List[Dict]) -> Tuple[str, bool]:
        """递归提取 JSON 卡片中的所有文本"""
        if not isinstance(message_array, list):
            return "", False

        _MAX_CARD_DEPTH = 20  # 防恶意深层嵌套卡片导致递归过深

        def collect_strings(obj, collected: set, depth: int = 0):
            if depth > _MAX_CARD_DEPTH:
                return
            if isinstance(obj, dict):
                for key, value in obj.items():
                    if key in ['url', 'jumpUrl', 'preview', 'icon', 'appid', 'uin', 'scene', 'token', 'ctime', 'width', 'height', 'forward', 'autoSize']:
                        continue
                    collect_strings(value, collected, depth + 1)
            elif isinstance(obj, list):
                for item in obj:
                    collect_strings(item, collected, depth + 1)
            elif isinstance(obj, str):
                if len(obj.strip()) > 1 and not obj.strip().isdigit():
                    collected.add(obj.strip())

        for msg in message_array:
            if msg.get("type") != "json":
                continue
            data = msg.get("data", {})
            json_str = data.get("data") or data.get("content") or data.get("text")
            if isinstance(json_str, dict):
                card_data = json_str
            elif isinstance(json_str, str):
                try:
                    card_data = json.loads(json_str)
                except json.JSONDecodeError:
                    continue
            else:
                if isinstance(data, dict):
                    card_data = data
                else:
                    continue
            collected = set()
            collect_strings(card_data, collected)
            if collected:
                card_text = "；".join(list(collected))
                return f"卡片内容：{card_text}", True
        return "", False

    # ========== 提取 @ 和纯文本 ==========
    def extract_mention_and_text(self, message_array: List[Dict], bot_qq: int) -> Tuple[str, bool]:
        """提取纯文本和是否@机器人"""
        if not isinstance(message_array, list):
            return "", False
        self_id = str(bot_qq)
        is_mentioned = False
        text_parts = []
        for msg in message_array:
            if msg.get("type") == "at":
                qq = str(msg.get("data", {}).get("qq", ""))
                if qq == self_id:
                    is_mentioned = True
            elif msg.get("type") == "text":
                text_parts.append(msg.get("data", {}).get("text", ""))
        return "".join(text_parts).strip(), is_mentioned
